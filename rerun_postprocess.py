#!/usr/bin/env python3
"""Re-run only the post-training steps of run_code.sh for already-trained runs.

run_code.sh does: train.py -> per-run summary XLSX -> results workbook update
-> augmented plot. If training succeeded but a later step failed (e.g. the
results/baseline_cost_estimation.xlsx was missing), retraining is unnecessary:
every setting needed is stored in the per-run summary XLSX ("Configuration"
sheet), so this script re-runs just:

  AUGMENT=True  -> update_augmented_cost_estimation.py + plot_augmented_cost_estimation.py
  AUGMENT=False -> update_baseline_cost_estimation.py

Runs that share a GROUP_RUN_TIME (N_RUNS > 1) are passed together, as
run_code.sh does. Runs whose training did not succeed are skipped -- those
need a real rerun of run_code.sh.

Usage (activate the project env first):
  python rerun_postprocess.py saved/summary_xlsx/20261001_191622_3287490_basketball_act_True.xlsx
  python rerun_postprocess.py saved/summary_xlsx/20261001_*_True.xlsx --dry-run
"""

from __future__ import annotations

import argparse
import shlex
import subprocess
import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
BASELINE_XLSX = SCRIPT_DIR / "results" / "baseline_cost_estimation.xlsx"
AUGMENTED_XLSX = SCRIPT_DIR / "results" / "augmented_cost_estimation.xlsx"
PLOT_DIR = SCRIPT_DIR / "results" / "augmented_plots"


def read_config(path: Path) -> Dict[str, str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Configuration" not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"{path}: no 'Configuration' sheet")
    config: Dict[str, str] = {}
    for row in workbook["Configuration"].iter_rows(min_row=2, values_only=True):
        if len(row) < 3 or row[0] not in ("run", "run_variable"):
            continue
        # run_variable rows hold the exact env values run_code.sh used (UPPER_CASE keys).
        config[str(row[1])] = "" if row[2] is None else str(row[2])
    workbook.close()
    return config


def run(cmd: List[str], dry_run: bool) -> int:
    print("+ " + shlex.join(cmd), flush=True)
    if dry_run:
        return 0
    return subprocess.run(cmd, cwd=SCRIPT_DIR).returncode


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("summary_xlsx", nargs="+", type=Path, help="Per-run summary XLSX files (not *_aggregate.xlsx)")
    parser.add_argument("--dry-run", action="store_true", help="Print the commands without running them")
    parser.add_argument("--no-plot", action="store_true", help="Skip the augmented plot step")
    args = parser.parse_args()

    groups: Dict[str, List[tuple]] = defaultdict(list)
    exit_code = 0
    for path in args.summary_xlsx:
        path = path.resolve()
        if path.name.endswith("_aggregate.xlsx"):
            continue
        try:
            cfg = read_config(path)
        except Exception as exc:  # noqa: BLE001 -- report and keep going
            print(f"SKIP {path}: {exc}", file=sys.stderr)
            exit_code = 1
            continue
        status = cfg.get("run_status", "")
        if status != "SUCCESS":
            print(f"SKIP {path.name}: training status is '{status}', rerun run_code.sh for this one", file=sys.stderr)
            exit_code = 1
            continue
        groups[cfg["GROUP_RUN_TIME"]].append((int(cfg.get("run_index", 1)), path, cfg))

    py = sys.executable
    for group_time, runs in sorted(groups.items()):
        runs.sort(key=lambda item: item[0])
        cfg = runs[0][2]
        files = [str(path) for _, path, _ in runs]
        requested = cfg["N_RUNS"]
        if len(files) != int(requested):
            print(f"NOTE {group_time}: {len(files)} of {requested} run files given", file=sys.stderr)
        print(f"\n=== {group_time}: {cfg['TEST_DB']} / {cfg['CARDINALITY_TYPE']} / AUGMENT={cfg['AUGMENT']} ===")

        common = [
            "--requested-runs", requested,
            "--test-db", cfg["TEST_DB"],
            "--time-stamp", group_time,
            "--epochs", cfg["EPOCHS"],
        ]
        aug_flags = [
            "--test-augment", cfg["TEST_AUGMENT"],
            "--augment-pooling", cfg["AUGMENT_POOLING"],
            "--augment-refinement", cfg["AUGMENT_REFINEMENT"],
            "--augment-coarse-layers", cfg["AUGMENT_COARSE_LAYERS"],
            "--augment-include-inv", cfg["AUGMENT_INCLUDE_INV"],
            "--augment-refine-ret", cfg["AUGMENT_REFINE_RET"],
            "--lambda-struct", cfg["LAMBDA_STRUCT"],
        ]

        if cfg["AUGMENT"].lower() == "false":
            rc = run([py, "update_baseline_cost_estimation.py", *common,
                      "--cardinality-type", cfg["CARDINALITY_TYPE"],
                      "--output-xlsx", str(BASELINE_XLSX), *files], args.dry_run)
            exit_code = exit_code or rc
            continue

        if not BASELINE_XLSX.exists():
            print(f"FAIL {group_time}: {BASELINE_XLSX} is missing -- copy it over first", file=sys.stderr)
            exit_code = 1
            continue

        rc = run([py, "update_augmented_cost_estimation.py", *common,
                  "--cardinality-type", cfg["CARDINALITY_TYPE"],
                  "--baseline-xlsx", str(BASELINE_XLSX),
                  "--output-xlsx", str(AUGMENTED_XLSX),
                  *aug_flags, *files], args.dry_run)
        exit_code = exit_code or rc

        if not args.no_plot:
            rc = run([py, "plot_augmented_cost_estimation.py", *common,
                      "--cardinality", cfg["CARDINALITY_TYPE"],
                      "--baseline-xlsx", str(BASELINE_XLSX),
                      "--output-dir", str(PLOT_DIR),
                      "--seed", cfg["SEED"],
                      "--augment", cfg["AUGMENT"],
                      *aug_flags, *files], args.dry_run)
            exit_code = exit_code or rc

    return exit_code


if __name__ == "__main__":
    sys.exit(main())
